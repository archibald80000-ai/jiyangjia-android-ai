from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote
from typing import Iterable, Protocol

import faiss
import numpy as np

from .answer_policy import classify_answer_scope


VALID_STATUSES = {"approved", "draft", "rejected"}
CUSTOMER_ALLOWED_STATUSES = {"approved"}
SAFE_TRANSFER_TEXT = "这项信息我暂时不能确认，请咨询现场工作人员。"
STOP_TOKENS = {"有没有", "有没", "没有", "可以", "怎么", "什么", "现在", "今天", "你们", "我们", "一下", "多少", "是不是", "能不能"}
MIN_KNOWLEDGE_SCORE = 0.6
MIN_KEYWORD_ONLY_SCORE = 0.7
MIN_VECTOR_ONLY_SCORE = 0.82


class EmbeddingLike(Protocol):
    name: str

    async def embed(self, texts: list[str], request_id: str) -> dict[str, object]: ...


@dataclass(frozen=True)
class KnowledgeDocument:
    doc_id: str
    title: str
    text: str
    status: str
    source_uri: str | None = None
    source_type: str = "manual"
    reviewed_by: str | None = None
    reviewed_at: str | None = None


@dataclass(frozen=True)
class ParsedKnowledgeDocument:
    doc_id: str
    title: str
    text: str
    status: str
    source_uri: str
    source_type: str
    reviewed_by: str | None = None
    reviewed_at: str | None = None


class SQLiteKnowledgeStore:
    def __init__(self, db_path: str = ":memory:", *, faiss_index_path: str | None = None) -> None:
        self._conn = sqlite3.connect(db_path, check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        self._faiss_index_path = Path(faiss_index_path) if faiss_index_path else None
        self._vector_index: faiss.IndexFlatIP | None = None
        self._vector_chunk_ids: list[str] = []
        self._vector_dimensions = 0
        self._init_schema()
        self._rebuild_vector_index()

    def _init_schema(self) -> None:
        self._conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS knowledge_documents (
                doc_id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                text TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('approved', 'draft', 'rejected')),
                source_uri TEXT,
                source_type TEXT NOT NULL DEFAULT 'manual',
                checksum TEXT NOT NULL,
                reviewed_by TEXT,
                reviewed_at TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS knowledge_chunks (
                chunk_id TEXT PRIMARY KEY,
                doc_id TEXT NOT NULL,
                chunk_index INTEGER NOT NULL,
                title TEXT NOT NULL,
                text TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('approved', 'draft', 'rejected')),
                token_count INTEGER NOT NULL,
                source_uri TEXT,
                checksum TEXT NOT NULL,
                FOREIGN KEY(doc_id) REFERENCES knowledge_documents(doc_id) ON DELETE CASCADE
            );

            CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_chunks_fts
            USING fts5(chunk_id UNINDEXED, title, text, tokenize='unicode61');

            CREATE TABLE IF NOT EXISTS knowledge_embeddings (
                chunk_id TEXT PRIMARY KEY,
                provider TEXT NOT NULL,
                model TEXT NOT NULL,
                dimensions INTEGER NOT NULL,
                vector_ref TEXT NOT NULL,
                vector_json TEXT NOT NULL,
                checksum TEXT NOT NULL,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(chunk_id) REFERENCES knowledge_chunks(chunk_id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS knowledge_ingestion_runs (
                run_id TEXT PRIMARY KEY,
                request_id TEXT NOT NULL,
                status TEXT NOT NULL,
                document_count INTEGER NOT NULL,
                chunk_count INTEGER NOT NULL,
                error TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            """
        )
        self._conn.commit()

    async def index_documents(self, documents: Iterable[KnowledgeDocument], embedding_provider: EmbeddingLike, request_id: str) -> dict[str, object]:
        docs = list(documents)
        run_id = f"ing_{uuid.uuid4().hex[:16]}"
        try:
            document_count, chunks = self._upsert_documents_and_chunks(docs)
            if chunks:
                embedding = await embedding_provider.embed([chunk["text"] for chunk in chunks], request_id=request_id)
                vectors = _coerce_vectors(embedding.get("vectors"), expected_count=len(chunks))
                dimensions = int(embedding.get("dimensions") or (len(vectors[0]) if vectors else 0))
                provider = str(embedding.get("provider") or embedding_provider.name)
                model = str(embedding.get("model") or provider)
                self._store_embeddings(chunks, vectors, dimensions, provider, model)
            self._record_run(run_id, request_id, "done", document_count, len(chunks), None)
            self._rebuild_vector_index()
            self._save_faiss_index()
            return {"run_id": run_id, "documents": document_count, "chunks": len(chunks), "status": self.status()}
        except Exception as exc:
            self._record_run(run_id, request_id, "failed", len(docs), 0, _safe_error(exc))
            raise

    def upsert_many(self, documents: Iterable[KnowledgeDocument]) -> int:
        document_count, _chunks = self._upsert_documents_and_chunks(list(documents))
        return document_count

    def list_documents(self, *, limit: int = 200) -> list[dict[str, object]]:
        rows = self._conn.execute(
            """
            SELECT doc_id, title, status, source_uri, source_type, reviewed_by,
                   reviewed_at, created_at, updated_at
            FROM knowledge_documents
            ORDER BY updated_at DESC, doc_id ASC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
        return [dict(row) for row in rows]

    def set_documents_status(
        self,
        doc_ids: Iterable[str],
        status: str,
        *,
        reviewed_by: str | None = None,
        reviewed_at: str | None = None,
    ) -> int:
        if status not in VALID_STATUSES:
            raise ValueError(f"invalid knowledge status: {status}")
        ids = [item for item in dict.fromkeys(doc_ids) if item]
        for doc_id in ids:
            self._conn.execute(
                """
                UPDATE knowledge_documents
                SET status = ?, reviewed_by = COALESCE(?, reviewed_by),
                    reviewed_at = COALESCE(?, reviewed_at), updated_at = CURRENT_TIMESTAMP
                WHERE doc_id = ?
                """,
                (status, reviewed_by, reviewed_at, doc_id),
            )
            self._conn.execute("UPDATE knowledge_chunks SET status = ? WHERE doc_id = ?", (status, doc_id))
        self._conn.commit()
        self._rebuild_vector_index()
        self._save_faiss_index()
        return len(ids)

    def _upsert_documents_and_chunks(self, documents: list[KnowledgeDocument]) -> tuple[int, list[dict[str, object]]]:
        chunks_for_embedding: list[dict[str, object]] = []
        for doc in documents:
            _validate_document(doc)
            checksum = _sha256_text(doc.text)
            self._conn.execute(
                """
                INSERT INTO knowledge_documents(doc_id, title, text, status, source_uri, source_type, checksum, reviewed_by, reviewed_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(doc_id) DO UPDATE SET
                    title=excluded.title,
                    text=excluded.text,
                    status=excluded.status,
                    source_uri=excluded.source_uri,
                    source_type=excluded.source_type,
                    checksum=excluded.checksum,
                    reviewed_by=excluded.reviewed_by,
                    reviewed_at=excluded.reviewed_at,
                    updated_at=CURRENT_TIMESTAMP
                """,
                (doc.doc_id, doc.title, doc.text, doc.status, doc.source_uri, doc.source_type, checksum, doc.reviewed_by, doc.reviewed_at),
            )
            self._delete_chunks_for_document(doc.doc_id)
            for chunk_index, text in enumerate(chunk_text(doc.text)):
                chunk_id = f"{doc.doc_id}::chunk-{chunk_index:04d}"
                chunk_checksum = _sha256_text(text)
                row = {
                    "chunk_id": chunk_id,
                    "doc_id": doc.doc_id,
                    "chunk_index": chunk_index,
                    "title": doc.title,
                    "text": text,
                    "status": doc.status,
                    "source_uri": doc.source_uri,
                    "checksum": chunk_checksum,
                }
                self._conn.execute(
                    """
                    INSERT INTO knowledge_chunks(chunk_id, doc_id, chunk_index, title, text, status, token_count, source_uri, checksum)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (chunk_id, doc.doc_id, chunk_index, doc.title, text, doc.status, _rough_token_count(text), doc.source_uri, chunk_checksum),
                )
                self._conn.execute(
                    "INSERT INTO knowledge_chunks_fts(chunk_id, title, text) VALUES (?, ?, ?)",
                    (chunk_id, doc.title, text),
                )
                if doc.status != "rejected":
                    chunks_for_embedding.append(row)
        self._conn.commit()
        return len(documents), chunks_for_embedding

    def _delete_chunks_for_document(self, doc_id: str) -> None:
        old = self._conn.execute("SELECT chunk_id FROM knowledge_chunks WHERE doc_id = ?", (doc_id,)).fetchall()
        chunk_ids = [row["chunk_id"] for row in old]
        for chunk_id in chunk_ids:
            self._conn.execute("DELETE FROM knowledge_chunks_fts WHERE chunk_id = ?", (chunk_id,))
            self._conn.execute("DELETE FROM knowledge_embeddings WHERE chunk_id = ?", (chunk_id,))
        self._conn.execute("DELETE FROM knowledge_chunks WHERE doc_id = ?", (doc_id,))

    def _store_embeddings(self, chunks: list[dict[str, object]], vectors: list[list[float]], dimensions: int, provider: str, model: str) -> None:
        for chunk, vector in zip(chunks, vectors):
            chunk_id = str(chunk["chunk_id"])
            self._conn.execute(
                """
                INSERT INTO knowledge_embeddings(chunk_id, provider, model, dimensions, vector_ref, vector_json, checksum)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(chunk_id) DO UPDATE SET
                    provider=excluded.provider,
                    model=excluded.model,
                    dimensions=excluded.dimensions,
                    vector_ref=excluded.vector_ref,
                    vector_json=excluded.vector_json,
                    checksum=excluded.checksum,
                    updated_at=CURRENT_TIMESTAMP
                """,
                (chunk_id, provider, model, dimensions, chunk_id, json.dumps(vector, separators=(",", ":")), str(chunk["checksum"])),
            )
        self._conn.commit()

    def _record_run(self, run_id: str, request_id: str, status: str, document_count: int, chunk_count: int, error: str | None) -> None:
        self._conn.execute(
            """
            INSERT INTO knowledge_ingestion_runs(run_id, request_id, status, document_count, chunk_count, error)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (run_id, request_id, status, document_count, chunk_count, error),
        )
        self._conn.commit()

    def status(self) -> dict[str, object]:
        rows = self._conn.execute("SELECT status, COUNT(*) AS count FROM knowledge_documents GROUP BY status").fetchall()
        documents = {"approved": 0, "draft": 0, "rejected": 0}
        documents.update({row["status"]: int(row["count"]) for row in rows})
        chunk_count = int(self._conn.execute("SELECT COUNT(*) AS count FROM knowledge_chunks").fetchone()["count"])
        embedding_count = int(self._conn.execute("SELECT COUNT(*) AS count FROM knowledge_embeddings").fetchone()["count"])
        latest_run = self._conn.execute(
            "SELECT run_id, status, document_count, chunk_count, error, created_at FROM knowledge_ingestion_runs ORDER BY created_at DESC LIMIT 1"
        ).fetchone()
        return {
            "documents": documents,
            "chunks": chunk_count,
            "embeddings": embedding_count,
            "faiss": {"ready": self._vector_index is not None and self._vector_index.ntotal > 0, "vectors": self._vector_index.ntotal if self._vector_index else 0, "dimensions": self._vector_dimensions},
            "latest_run": dict(latest_run) if latest_run else None,
        }

    def classify_query(self, query: str) -> dict[str, object]:
        scope = classify_answer_scope(query)
        return {"action": "search", "scope": scope, "category": None, "message": None}

    async def search(
        self,
        query: str,
        embedding_provider: EmbeddingLike | None = None,
        request_id: str = "knowledge-search",
        top_k: int = 3,
        include_draft: bool = False,
    ) -> list[dict[str, object]]:
        policy = self.classify_query(query)
        if policy["action"] != "search":
            return []
        if policy["scope"] == "general":
            return []
        # Search drafts as a shadow corpus even for customer requests. If the
        # strongest relevant document is still draft, the information is not
        # approved for use and tangential approved documents must not mask it.
        candidate_statuses = ("approved", "draft")
        vector_scores: dict[str, float] = {}
        if embedding_provider is not None and self._vector_index is not None and self._vector_index.ntotal > 0 and self._vector_dimensions > 1:
            vector_scores = await self._vector_search(query, embedding_provider, request_id, top_k=max(top_k * 4, 8), allowed=candidate_statuses)
        keyword_scores = self._keyword_search(query, top_k=max(top_k * 4, 8), allowed=candidate_statuses)
        merged = self._merge_scores(vector_scores, keyword_scores, allowed=candidate_statuses)
        if include_draft:
            return merged[:top_k]
        if merged and merged[0]["status"] == "draft":
            return []
        return [match for match in merged if match["status"] == "approved"][:top_k]

    async def _vector_search(self, query: str, embedding_provider: EmbeddingLike, request_id: str, top_k: int, allowed: tuple[str, ...]) -> dict[str, float]:
        embedding = await embedding_provider.embed([query], request_id=request_id)
        vectors = _coerce_vectors(embedding.get("vectors"), expected_count=1)
        query_vector = _normalize_vectors(np.array(vectors, dtype="float32"))
        if query_vector.shape[1] != self._vector_dimensions:
            return {}
        distances, indices = self._vector_index.search(query_vector, min(top_k, self._vector_index.ntotal))
        scores: dict[str, float] = {}
        for score, index in zip(distances[0], indices[0]):
            if index < 0 or index >= len(self._vector_chunk_ids):
                continue
            chunk_id = self._vector_chunk_ids[int(index)]
            row = self._chunk_by_id(chunk_id)
            if row and row["status"] in allowed and float(score) >= 0.25:
                scores[chunk_id] = max(scores.get(chunk_id, 0.0), float(score))
        return scores

    def _keyword_search(self, query: str, top_k: int, allowed: tuple[str, ...]) -> dict[str, float]:
        tokens = _tokens(query)
        if not tokens:
            return {}
        rows = self._candidate_chunks(allowed)
        scores: dict[str, float] = {}
        for row in rows:
            haystack = f"{row['title']} {row['text']}".lower()
            hits = sum(1 for token in tokens if token in haystack)
            has_strong_token = any(len(token) >= 3 and token in haystack for token in tokens)
            if hits >= 2 or has_strong_token:
                scores[row["chunk_id"]] = min(1.0, 0.25 + hits / max(4, len(tokens)))
        return dict(sorted(scores.items(), key=lambda item: -item[1])[:top_k])

    def _merge_scores(self, vector_scores: dict[str, float], keyword_scores: dict[str, float], allowed: tuple[str, ...]) -> list[dict[str, object]]:
        chunk_ids = set(vector_scores) | set(keyword_scores)
        results: list[dict[str, object]] = []
        for chunk_id in chunk_ids:
            row = self._chunk_by_id(chunk_id)
            if not row or row["status"] not in allowed:
                continue
            vector_score = vector_scores.get(chunk_id, 0.0)
            keyword_score = keyword_scores.get(chunk_id, 0.0)
            if keyword_score > 0.0:
                # Keep ranking stable when a chunk falls just outside the
                # bounded FAISS candidate window: lexical evidence contributes
                # the same base weight whether or not a vector score is present.
                combined = keyword_score * 0.8 + max(vector_score, 0.0) * 0.2
            else:
                combined = max(vector_score, 0.0) * 0.75
            strong_keyword_evidence = keyword_score >= MIN_KEYWORD_ONLY_SCORE
            if combined < MIN_KNOWLEDGE_SCORE and not strong_keyword_evidence:
                continue
            if keyword_score == 0.0 and vector_score < MIN_VECTOR_ONLY_SCORE:
                continue
            confidence = min(0.99, max(0.05, combined))
            results.append(_match_from_row(row, confidence=confidence, vector_score=vector_score, keyword_score=keyword_score))
        results.sort(key=lambda item: (-float(item["confidence"]), str(item["id"])))
        return results

    def _candidate_chunks(self, allowed: tuple[str, ...]) -> list[sqlite3.Row]:
        placeholders = ",".join("?" for _ in allowed)
        return self._conn.execute(f"SELECT * FROM knowledge_chunks WHERE status IN ({placeholders})", allowed).fetchall()

    def _chunk_by_id(self, chunk_id: str) -> sqlite3.Row | None:
        return self._conn.execute("SELECT * FROM knowledge_chunks WHERE chunk_id = ?", (chunk_id,)).fetchone()

    def _rebuild_vector_index(self) -> None:
        rows = self._conn.execute(
            """
            SELECT e.chunk_id, e.vector_json, e.dimensions
            FROM knowledge_embeddings e
            JOIN knowledge_chunks c ON c.chunk_id = e.chunk_id
            WHERE c.status IN ('approved', 'draft')
            ORDER BY e.chunk_id
            """
        ).fetchall()
        if not rows:
            self._vector_index = None
            self._vector_chunk_ids = []
            self._vector_dimensions = 0
            return
        dimensions = int(rows[0]["dimensions"])
        vectors = []
        chunk_ids = []
        for row in rows:
            if int(row["dimensions"]) != dimensions:
                continue
            vectors.append(json.loads(row["vector_json"]))
            chunk_ids.append(row["chunk_id"])
        if not vectors:
            self._vector_index = None
            self._vector_chunk_ids = []
            self._vector_dimensions = 0
            return
        matrix = _normalize_vectors(np.array(vectors, dtype="float32"))
        index = faiss.IndexFlatIP(matrix.shape[1])
        index.add(matrix)
        self._vector_index = index
        self._vector_chunk_ids = chunk_ids
        self._vector_dimensions = matrix.shape[1]

    def _save_faiss_index(self) -> None:
        if self._faiss_index_path is None:
            return
        if self._vector_index is None:
            self._faiss_index_path.unlink(missing_ok=True)
            return
        self._faiss_index_path.parent.mkdir(parents=True, exist_ok=True)
        faiss.write_index(self._vector_index, str(self._faiss_index_path))


def parse_selected_knowledge_path(path: Path, *, default_status: str = "draft") -> list[KnowledgeDocument]:
    resolved = Path(path)
    if resolved.is_dir():
        documents: list[KnowledgeDocument] = []
        for child in sorted(resolved.iterdir()):
            if child.name.startswith("."):
                continue
            if child.suffix.lower() in {".json", ".yaml", ".yml", ".md", ".txt", ".pdf", ".docx"}:
                documents.extend(parse_selected_knowledge_path(child, default_status=default_status))
        return documents
    suffix = resolved.suffix.lower()
    if suffix == ".json":
        return _parse_json_documents(resolved, default_status=default_status)
    if suffix in {".yaml", ".yml"}:
        return _parse_yaml_documents(resolved, default_status=default_status)
    if suffix in {".md", ".txt"}:
        text = resolved.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        text = _read_pdf_text(resolved)
    elif suffix == ".docx":
        text = _read_docx_text(resolved)
    else:
        raise ValueError(f"unsupported knowledge file: {resolved}")
    return [
        KnowledgeDocument(
            doc_id=_doc_id_from_path(resolved),
            title=_title_from_text(text, resolved.stem),
            text=text,
            status=default_status,
            source_uri=resolved.as_posix(),
            source_type=suffix.lstrip("."),
        )
    ]


def chunk_text(text: str, *, max_chars: int = 700, overlap: int = 80) -> list[str]:
    normalized = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not normalized:
        return []
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start : start + max_chars].strip())
                start += max_chars - overlap
            continue
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) > max_chars and current:
            chunks.append(current.strip())
            current = paragraph
        else:
            current = candidate
    if current:
        chunks.append(current.strip())
    return chunks


def _parse_json_documents(path: Path, *, default_status: str) -> list[KnowledgeDocument]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    items = payload.get("documents") or payload.get("items") if isinstance(payload, dict) else payload
    if not isinstance(items, list):
        raise ValueError(f"json knowledge file must contain documents/items list: {path}")
    documents = []
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("knowledge item must be an object")
        status = _normalize_status(str(item.get("status") or default_status))
        answer = str(item.get("answer") or item.get("text") or "").strip()
        questions = item.get("questions") if isinstance(item.get("questions"), list) else []
        text = "\n".join([*(str(q) for q in questions), answer]).strip()
        documents.append(
            KnowledgeDocument(
                doc_id=str(item.get("id") or _doc_id_from_path(path)),
                title=str(item.get("title") or (questions[0] if questions else item.get("id") or path.stem)),
                text=text,
                status=status,
                source_uri=str(item.get("source_uri") or item.get("source") or path.as_posix()),
                source_type=str(item.get("source_type") or "faq"),
                reviewed_by=item.get("reviewed_by"),
                reviewed_at=item.get("reviewed_at"),
            )
        )
    return documents


def _parse_yaml_documents(path: Path, *, default_status: str) -> list[KnowledgeDocument]:
    import yaml

    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    items = payload.get("documents") or payload.get("items") if isinstance(payload, dict) else payload
    if not isinstance(items, list):
        raise ValueError(f"yaml knowledge file must contain documents/items list: {path}")
    documents = []
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("knowledge item must be an object")
        status = _normalize_status(str(item.get("status") or default_status))
        answer = str(item.get("answer") or item.get("text") or "").strip()
        questions = item.get("questions") if isinstance(item.get("questions"), list) else []
        text = "\n".join([*(str(q) for q in questions), answer]).strip()
        documents.append(
            KnowledgeDocument(
                doc_id=str(item.get("id") or _doc_id_from_path(path)),
                title=str(item.get("title") or (questions[0] if questions else item.get("id") or path.stem)),
                text=text,
                status=status,
                source_uri=str(item.get("source_uri") or item.get("source") or path.as_posix()),
                source_type=str(item.get("source_type") or "faq"),
                reviewed_by=item.get("reviewed_by"),
                reviewed_at=item.get("reviewed_at"),
            )
        )
    return documents


def _read_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _read_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _normalize_status(value: str) -> str:
    normalized = value.strip().lower()
    if normalized == "mock":
        return "draft"
    if normalized not in VALID_STATUSES:
        raise ValueError(f"invalid knowledge status: {value}")
    return normalized


def _validate_document(doc: KnowledgeDocument) -> None:
    if doc.status not in VALID_STATUSES:
        raise ValueError(f"invalid knowledge status: {doc.status}")
    if not doc.doc_id.strip() or not doc.title.strip() or not doc.text.strip():
        raise ValueError("knowledge document requires id, title and text")


def _coerce_vectors(value: object, *, expected_count: int) -> list[list[float]]:
    if not isinstance(value, list) or len(value) < expected_count:
        raise ValueError("embedding result did not contain enough vectors")
    vectors: list[list[float]] = []
    for vector in value[:expected_count]:
        if not isinstance(vector, list) or not vector:
            raise ValueError("embedding vector must be a non-empty list")
        vectors.append([float(item) for item in vector])
    dimensions = {len(vector) for vector in vectors}
    if len(dimensions) != 1:
        raise ValueError("embedding vectors must have consistent dimensions")
    return vectors


def _normalize_vectors(matrix: np.ndarray) -> np.ndarray:
    norms = np.linalg.norm(matrix, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return matrix / norms


def _match_from_row(row: sqlite3.Row, *, confidence: float, vector_score: float, keyword_score: float) -> dict[str, object]:
    return {
        "id": row["doc_id"],
        "chunk_id": row["chunk_id"],
        "title": row["title"],
        "status": row["status"],
        "score": round(confidence, 4),
        "confidence": round(confidence, 4),
        "vector_score": round(float(vector_score), 4),
        "keyword_score": round(float(keyword_score), 4),
        "source": {"uri": _public_source_uri(row["source_uri"], row["doc_id"]), "chunk_id": row["chunk_id"]},
        "excerpt": row["text"][:240],
    }


def _public_source_uri(source_uri: str | None, doc_id: str) -> str:
    raw = str(source_uri or "").strip()
    scheme = re.match(r"^([a-z][a-z0-9+.-]*)://", raw, flags=re.IGNORECASE)
    if scheme and scheme.group(1).lower() in {"http", "https", "admin", "manual", "knowledge"}:
        return raw
    return f"knowledge://{quote(str(doc_id), safe='-._~')}"


def _tokens(text: str) -> list[str]:
    normalized = re.sub(r"\s+", " ", text.strip().lower())
    # The brand selects the business corpus but should not dilute lexical
    # relevance inside that corpus. Boundary documents intentionally omit the
    # repeated brand prefix, so score only the customer's substantive terms.
    normalized = normalized.replace("积养家", " ")
    if not normalized:
        return []
    ascii_tokens = re.findall(r"[a-z0-9_]+", normalized)
    chinese_chunks = re.findall(r"[\u4e00-\u9fff]{2,}", normalized)
    bigrams: list[str] = []
    trigrams: list[str] = []
    for chunk in chinese_chunks:
        bigrams.extend(chunk[i : i + 2] for i in range(max(0, len(chunk) - 1)))
        trigrams.extend(chunk[i : i + 3] for i in range(max(0, len(chunk) - 2)))
    tokens = ascii_tokens + chinese_chunks + trigrams + bigrams
    return [token for token in dict.fromkeys(tokens) if token not in STOP_TOKENS]


def _rough_token_count(text: str) -> int:
    return max(1, len(_tokens(text)) or len(text) // 2)


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest().upper()


def _doc_id_from_path(path: Path) -> str:
    digest = hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:12]
    return f"doc_{path.stem}_{digest}"


def _title_from_text(text: str, fallback: str) -> str:
    for line in text.splitlines():
        stripped = line.strip().lstrip("#").strip()
        if stripped:
            return stripped[:120]
    return fallback


def _safe_error(exc: Exception) -> str:
    text = str(exc)
    if len(text) > 240:
        return text[:240] + "..."
    return text
