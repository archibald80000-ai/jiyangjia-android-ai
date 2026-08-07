from __future__ import annotations

import uuid
from pathlib import Path
from io import BytesIO

from fastapi import FastAPI, Request
from fastapi.testclient import TestClient

from gateway.app.admin_routes import create_admin_router
from gateway.app.admin_store import AdminContentStore, DISPLAY_DEFAULT_MIGRATION_ID
from gateway.app.admin_ui import admin_page
from gateway.app.config import Settings
from gateway.app.knowledge import SQLiteKnowledgeStore
from gateway.app.providers import MockEmbeddingProvider


def _client(
    tmp_path: Path,
    *,
    app_env: str = "development",
    admin_token: str = "",
    embedding_provider=None,
) -> tuple[TestClient, AdminContentStore, SQLiteKnowledgeStore]:
    settings = Settings(
        app_env=app_env,
        admin_token=admin_token,
        admin_db_path=str(tmp_path / "admin" / "admin.db"),
        admin_upload_dir=str(tmp_path / "knowledge" / "uploads"),
        asset_dir=str(tmp_path / "assets"),
        knowledge_db_path=str(tmp_path / "knowledge" / "knowledge.db"),
        knowledge_faiss_path=str(tmp_path / "knowledge" / "faiss.index"),
    )
    Path(settings.knowledge_db_path).parent.mkdir(parents=True, exist_ok=True)
    knowledge = SQLiteKnowledgeStore(settings.knowledge_db_path, faiss_index_path=settings.knowledge_faiss_path)
    admin = AdminContentStore(settings.admin_db_path, upload_dir=settings.admin_upload_dir, asset_dir=settings.asset_dir)
    app = FastAPI()

    @app.middleware("http")
    async def request_id(request: Request, call_next):
        request.state.request_id = request.headers.get("X-Request-Id") or str(uuid.uuid4())
        return await call_next(request)

    app.include_router(
        create_admin_router(
            settings=settings,
            knowledge_store=knowledge,
            admin_store=admin,
            embedding_provider=embedding_provider or MockEmbeddingProvider(),
            provider_status=lambda: {
                name: {"provider": "mock", "ready": True, "missing": []}
                for name in ("asr", "tts", "llm", "embedding")
            },
            embedding_error=lambda: None,
        )
    )
    return TestClient(app), admin, knowledge


def test_four_admin_pages_and_system_status_load(tmp_path: Path) -> None:
    client, _admin, _knowledge = _client(tmp_path)
    for section in ("system", "knowledge", "avatar", "display"):
        response = client.get(f"/admin/{section}")
        assert response.status_code == 200
        assert "正在加载" in response.text

    status = client.get("/api/v1/admin/system/status")
    assert status.status_code == 200
    payload = status.json()
    assert payload["gateway"]["status"] == "healthy"
    assert payload["content"]["display_profiles"] == 4
    assert payload["admin_auth"]["configured"] is False


def test_admin_ui_action_dispatch_contract_and_complete_display_fields() -> None:
    knowledge_html = admin_page("knowledge")
    display_html = admin_page("display")
    assert "const [verb,id]=command.split(':'); const [kind,op]=verb.split('-');" in knowledge_html
    for command in ("knowledge-preview:", "knowledge-approve:", "knowledge-publish:", "knowledge-reject:"):
        assert command in knowledge_html
    for field in (
        "character_anchor_x",
        "character_anchor_y",
        "character_scale",
        "safe_left",
        "safe_right",
        "safe_bottom",
        "subtitle_font_px",
        "button_x",
        "button_y",
        "avatar_id",
        "background_id",
    ):
        assert f'name="{field}"' in display_html
    assert 'name="width_px" type="number" min="320" max="7680" value="1080"' in display_html
    assert 'name="height_px" type="number" min="320" max="7680" value="1920"' in display_html
    assert '<option value="portrait">竖屏</option><option value="landscape">横屏</option>' in display_html


def test_knowledge_requires_publish_before_customer_search(tmp_path: Path) -> None:
    client, _admin, knowledge = _client(tmp_path)
    uploaded = client.post(
        "/api/v1/admin/knowledge/upload",
        files={"file": ("store-guide.md", "# 到店指引\n\n积养家门店到店后请先咨询前台。".encode("utf-8"), "text/markdown")},
    )
    assert uploaded.status_code == 200
    run_id = uploaded.json()["run"]["run_id"]
    assert uploaded.json()["run"]["status"] == "parsed"
    assert knowledge.status()["documents"]["draft"] == 1

    assert client.post(f"/api/v1/admin/knowledge/{run_id}/preview").json()["preview"][0]["source"].startswith("admin://knowledge/")
    assert client.post(f"/api/v1/admin/knowledge/{run_id}/approve").json()["run"]["status"] == "approved"
    assert _run(knowledge.search("到店指引", MockEmbeddingProvider(), "req-before-publish")) == []
    assert knowledge.status()["documents"]["approved"] == 0

    published = client.post(f"/api/v1/admin/knowledge/{run_id}/publish")
    assert published.status_code == 200
    assert published.json()["run"]["status"] == "published"
    matches = _run(knowledge.search("到店指引", MockEmbeddingProvider(), "req-admin-search"))
    assert matches
    assert matches[0]["status"] == "approved"
    assert matches[0]["source"]["uri"].startswith("admin://knowledge/")


def test_rejected_upload_never_enters_customer_search(tmp_path: Path) -> None:
    client, _admin, knowledge = _client(tmp_path)
    uploaded = client.post(
        "/api/v1/admin/knowledge/upload",
        files={"file": ("draft.txt", "待确认的内部服务说明".encode("utf-8"), "text/plain")},
    )
    run_id = uploaded.json()["run"]["run_id"]
    rejected = client.post(f"/api/v1/admin/knowledge/{run_id}/reject")
    assert rejected.status_code == 200
    assert knowledge.status()["documents"]["rejected"] == 1
    assert _run(knowledge.search("内部服务说明", MockEmbeddingProvider(), "req-rejected")) == []


def test_failed_embedding_publish_returns_document_to_draft(tmp_path: Path) -> None:
    class FailingEmbeddingProvider:
        name = "failing"

        async def embed(self, texts, request_id):
            raise RuntimeError("controlled embedding failure")

    client, admin, knowledge = _client(tmp_path, embedding_provider=FailingEmbeddingProvider())
    uploaded = client.post(
        "/api/v1/admin/knowledge/upload",
        files={"file": ("failure.md", b"# Publish failure\n\nCustomer invisible content.", "text/markdown")},
    )
    run_id = uploaded.json()["run"]["run_id"]
    assert client.post(f"/api/v1/admin/knowledge/{run_id}/approve").status_code == 200
    failed = client.post(f"/api/v1/admin/knowledge/{run_id}/publish")
    assert failed.status_code == 502
    assert admin.get_knowledge_run(run_id)["status"] == "failed"
    assert knowledge.status()["documents"]["approved"] == 0
    assert knowledge.status()["documents"]["draft"] == 1
    assert _run(knowledge.search("Publish failure", MockEmbeddingProvider(), "req-failed")) == []


def test_pdf_docx_markdown_and_text_upload_parse(tmp_path: Path) -> None:
    from docx import Document

    client, _admin, _knowledge = _client(tmp_path)
    docx_buffer = BytesIO()
    document = Document()
    document.add_paragraph("DOCX 服务指引")
    document.add_paragraph("到店后请咨询现场工作人员。")
    document.save(docx_buffer)
    cases = (
        ("guide.pdf", _simple_pdf("PDF service guide"), "application/pdf"),
        ("guide.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("guide.md", b"# Markdown guide\n\nConfirmed content.", "text/markdown"),
        ("guide.txt", b"TXT guide\n\nConfirmed content.", "text/plain"),
    )
    for filename, content, content_type in cases:
        response = client.post("/api/v1/admin/knowledge/upload", files={"file": (filename, content, content_type)})
        assert response.status_code == 200, response.text
        assert response.json()["run"]["source_type"] == Path(filename).suffix.lstrip(".")
        assert response.json()["run"]["chunk_count"] >= 1


def test_asset_publish_manifest_rollback_delete_and_persistence(tmp_path: Path) -> None:
    client, admin, _knowledge = _client(tmp_path)
    first = client.post(
        "/api/v1/admin/avatar",
        data={"name": "门店待机 A", "version": "v1", "asset_type": "video"},
        files={"file": ("idle-v1.mp4", _mp4_bytes(b"v1"), "video/mp4")},
    ).json()["asset"]
    second = client.post(
        "/api/v1/admin/avatar",
        data={"name": "门店待机 B", "version": "v2", "asset_type": "video"},
        files={"file": ("idle-v2.mp4", _mp4_bytes(b"v2"), "video/mp4")},
    ).json()["asset"]
    image = client.post(
        "/api/v1/admin/avatar",
        data={"name": "门店背景", "version": "v1", "asset_type": "image"},
        files={"file": ("background.png", _png_bytes(), "image/png")},
    ).json()["asset"]

    assert client.get(f"/api/v1/assets/{first['avatar_id']}").status_code == 404
    assert client.get(f"/api/v1/admin/avatar/{first['avatar_id']}/file").content == _mp4_bytes(b"v1")

    assert client.post(f"/api/v1/admin/avatar/{first['avatar_id']}/publish").status_code == 200
    assert client.post(f"/api/v1/admin/avatar/{second['avatar_id']}/publish").json()["manifest"]["video"]["version"] == "v2"
    assert client.post(f"/api/v1/admin/avatar/{first['avatar_id']}/rollback").json()["manifest"]["video"]["version"] == "v1"
    assert client.post(f"/api/v1/admin/avatar/{image['avatar_id']}/publish").status_code == 200
    assert client.get(f"/api/v1/assets/{first['avatar_id']}").content == _mp4_bytes(b"v1")
    assert client.delete(f"/api/v1/admin/avatar/{second['avatar_id']}/delete").status_code == 200
    manifest = client.get("/api/v1/assets/manifest").json()
    assert manifest["video"]["sha256"] == first["sha256"]
    assert manifest["background"]["version"] == "v1"

    reopened = AdminContentStore(str(tmp_path / "admin" / "admin.db"), upload_dir=str(tmp_path / "knowledge" / "uploads"), asset_dir=str(tmp_path / "assets"))
    assert reopened.counts()["assets"] == 2
    assert reopened.manifest()["video"]["version"] == "v1"
    assert (admin.asset_dir / "manifest.json").is_file()


def test_display_presets_custom_match_and_persistence(tmp_path: Path) -> None:
    client, _admin, _knowledge = _client(tmp_path)
    profiles = client.get("/api/v1/admin/display").json()["profiles"]
    assert {(p["width_px"], p["height_px"]) for p in profiles} >= {(1920, 1080), (3840, 2160), (1280, 720), (1080, 1920)}
    default_profile = next(p for p in profiles if p["is_default"])
    assert (default_profile["width_px"], default_profile["height_px"], default_profile["orientation"]) == (1080, 1920, "portrait")
    assert client.get("/api/v1/display/profile").json()["profile"]["profile_id"] == "display-1080x1920"

    custom_payload = {
        "profile_name": "门店定制 1600x900",
        "width_px": 1600,
        "height_px": 900,
        "orientation": "landscape",
        "scale_mode": "crop",
        "character_anchor_x": 0.62,
        "character_anchor_y": 0.5,
        "character_scale": 1.1,
        "subtitle_safe_area": {"left": 0.1, "right": 0.1, "bottom": 0.12},
        "subtitle_font_px": 44,
        "button_positions": {"consult": {"x": 0.5, "y": 0.9}},
        "status": "active",
    }
    created = client.post("/api/v1/admin/display", json=custom_payload)
    assert created.status_code == 200
    profile_id = created.json()["profile"]["profile_id"]
    assert client.post(f"/api/v1/admin/display/{profile_id}/set-default").status_code == 200
    matched = client.get("/api/v1/display/profile?width=1600&height=900&orientation=landscape").json()["profile"]
    assert matched["profile_id"] == profile_id
    assert matched["scale_mode"] == "crop"

    reopened = AdminContentStore(str(tmp_path / "admin" / "admin.db"), upload_dir=str(tmp_path / "knowledge" / "uploads"), asset_dir=str(tmp_path / "assets"))
    assert reopened.get_display_profile(profile_id)["is_default"] is True
    assert reopened.default_display_profile()["profile_id"] == profile_id


def test_existing_admin_database_migrates_to_portrait_default_once(tmp_path: Path) -> None:
    _client_instance, admin, _knowledge = _client(tmp_path)
    assert admin.set_default_profile("display-1920x1080")["is_default"] is True
    admin._conn.execute("DELETE FROM admin_migrations WHERE migration_id = ?", (DISPLAY_DEFAULT_MIGRATION_ID,))
    admin._conn.commit()

    reopened = AdminContentStore(
        str(tmp_path / "admin" / "admin.db"),
        upload_dir=str(tmp_path / "knowledge" / "uploads"),
        asset_dir=str(tmp_path / "assets"),
    )
    assert reopened.default_display_profile()["profile_id"] == "display-1080x1920"
    assert reopened.set_default_profile("display-1920x1080")["is_default"] is True

    reopened_again = AdminContentStore(
        str(tmp_path / "admin" / "admin.db"),
        upload_dir=str(tmp_path / "knowledge" / "uploads"),
        asset_dir=str(tmp_path / "assets"),
    )
    assert reopened_again.default_display_profile()["profile_id"] == "display-1920x1080"


def test_asset_content_signature_and_display_binding_are_enforced(tmp_path: Path) -> None:
    client, _admin, _knowledge = _client(tmp_path)
    invalid = client.post(
        "/api/v1/admin/avatar",
        data={"name": "Invalid", "version": "v1", "asset_type": "video"},
        files={"file": ("invalid.mp4", b"not-a-video", "video/mp4")},
    )
    assert invalid.status_code == 422
    assert invalid.json()["detail"]["code"] == "INVALID_ASSET_CONTENT"

    draft = client.post(
        "/api/v1/admin/avatar",
        data={"name": "Draft", "version": "v1", "asset_type": "video"},
        files={"file": ("draft.mp4", _mp4_bytes(b"draft"), "video/mp4")},
    ).json()["asset"]
    payload = _display_payload("Bound display") | {"avatar_id": draft["avatar_id"]}
    blocked = client.post("/api/v1/admin/display", json=payload)
    assert blocked.status_code == 422
    assert blocked.json()["detail"]["code"] == "DISPLAY_ASSET_NOT_PUBLISHED"

    assert client.post(f"/api/v1/admin/avatar/{draft['avatar_id']}/publish").status_code == 200
    assert client.post("/api/v1/admin/display", json=payload).status_code == 200


def test_asset_manifest_etag_and_metadata_are_stable(tmp_path: Path) -> None:
    client, _admin, _knowledge = _client(tmp_path)
    asset = client.post(
        "/api/v1/admin/avatar",
        data={"name": "Sync video", "version": "v1", "asset_type": "video"},
        files={"file": ("sync.mp4", _mp4_bytes(b"sync"), "video/mp4")},
    ).json()["asset"]
    client.post(f"/api/v1/admin/avatar/{asset['avatar_id']}/publish")

    response = client.get("/api/v1/assets/manifest")
    assert response.status_code == 200
    assert response.headers["etag"].startswith('"')
    payload = response.json()
    assert payload["video"]["size_bytes"] == len(_mp4_bytes(b"sync"))
    assert payload["video"]["url"].endswith(asset["avatar_id"])

    unchanged = client.get(
        "/api/v1/assets/manifest",
        headers={"If-None-Match": response.headers["etag"]},
    )
    assert unchanged.status_code == 304
    assert unchanged.content == b""

    downloaded = client.get(payload["video"]["url"])
    assert downloaded.headers["etag"] == f'"{asset["sha256"]}"'
    assert downloaded.headers["cache-control"].endswith("immutable")


def test_production_admin_api_requires_configured_token(tmp_path: Path) -> None:
    blocked, _admin, _knowledge = _client(tmp_path / "missing", app_env="production")
    assert blocked.get("/api/v1/admin/system/status").status_code == 503

    protected, _admin, _knowledge = _client(tmp_path / "protected", app_env="production", admin_token="test-only-token")
    assert protected.get("/api/v1/admin/system/status").status_code == 401
    allowed = protected.get("/api/v1/admin/system/status", headers={"X-Admin-Token": "test-only-token"})
    assert allowed.status_code == 200
    assert "test-only-token" not in allowed.text


def _run(coroutine):
    import asyncio

    return asyncio.run(coroutine)


def _simple_pdf(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(text) + 34} >>\nstream\nBT /F1 12 Tf 72 720 Td ({text}) Tj ET\nendstream".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    return bytes(output)


def _mp4_bytes(marker: bytes) -> bytes:
    return b"\x00\x00\x00\x18ftypisom\x00\x00\x02\x00isomiso2" + marker


def _png_bytes() -> bytes:
    return b"\x89PNG\r\n\x1a\n" + b"test-png-content"


def _display_payload(name: str) -> dict[str, object]:
    return {
        "profile_name": name,
        "width_px": 1920,
        "height_px": 1080,
        "orientation": "landscape",
        "scale_mode": "fit",
        "character_anchor_x": 0.5,
        "character_anchor_y": 0.5,
        "character_scale": 1,
        "subtitle_safe_area": {"left": 0.08, "right": 0.08, "bottom": 0.08},
        "subtitle_font_px": 52,
        "button_positions": {"consult": {"x": 0.5, "y": 0.88}},
        "status": "active",
    }
