from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from gateway.app.knowledge import KnowledgeDocument, SQLiteKnowledgeStore, _expand_search_query, _public_source_uri, chunk_text, parse_selected_knowledge_path


class SemanticEmbeddingProvider:
    name = "semantic-test"

    async def embed(self, texts: list[str], request_id: str) -> dict[str, object]:
        vectors = [_vector_for_text(text) for text in texts]
        return {"provider": self.name, "model": "semantic-test-v1", "vectors": vectors, "dimensions": 4}


def test_schema_rejects_invalid_status() -> None:
    store = SQLiteKnowledgeStore(":memory:")
    with pytest.raises(ValueError):
        store.upsert_many([KnowledgeDocument("bad", "bad", "bad", "mock")])


def test_public_source_uri_preserves_public_schemes_and_redacts_local_paths() -> None:
    assert _public_source_uri("manual://faq/brand", "faq-brand") == "manual://faq/brand"
    assert _public_source_uri("admin://knowledge/run/chunk", "faq-brand") == "admin://knowledge/run/chunk"
    assert _public_source_uri("https://example.com/brand", "faq-brand") == "https://example.com/brand"
    assert _public_source_uri(r"E:\private\brand.md", "faq-brand") == "knowledge://faq-brand"
    assert _public_source_uri("file:///private/brand.md", "faq-brand") == "knowledge://faq-brand"


def test_chunk_text_splits_long_content() -> None:
    text = "标题\n\n" + "服务介绍" * 300
    chunks = chunk_text(text, max_chars=120, overlap=20)
    assert len(chunks) > 1
    assert all(len(chunk) <= 120 for chunk in chunks)


def test_index_builds_sqlite_metadata_and_faiss_search(tmp_path: Path) -> None:
    asyncio.run(_assert_index_builds_sqlite_metadata_and_faiss_search(tmp_path))


async def _assert_index_builds_sqlite_metadata_and_faiss_search(tmp_path: Path) -> None:
    store = SQLiteKnowledgeStore(":memory:", faiss_index_path=str(tmp_path / "faiss.index"))
    result = await store.index_documents(_documents(), SemanticEmbeddingProvider(), "req-index")

    assert result["documents"] == 4
    status = store.status()
    assert status["documents"]["approved"] == 2
    assert status["documents"]["draft"] == 1
    assert status["documents"]["rejected"] == 1
    assert status["chunks"] == 4
    assert status["embeddings"] == 3
    assert status["faiss"]["ready"] is True
    assert status["faiss"]["dimensions"] == 4
    assert (tmp_path / "faiss.index").exists()

    matches = await store.search("服务时间是什么", SemanticEmbeddingProvider(), "req-search", top_k=2)
    assert matches
    assert matches[0]["id"] == "approved_time"
    assert matches[0]["source"]["uri"] == "manual://time"
    assert matches[0]["chunk_id"].startswith("approved_time::chunk-")


def test_search_excludes_draft_and_rejected_by_default() -> None:
    asyncio.run(_assert_search_excludes_draft_and_rejected_by_default())


async def _assert_search_excludes_draft_and_rejected_by_default() -> None:
    store = SQLiteKnowledgeStore(":memory:")
    await store.index_documents(_documents(), SemanticEmbeddingProvider(), "req-index")

    default_matches = await store.search("服务包", SemanticEmbeddingProvider(), "req-draft", top_k=3)
    assert default_matches == []

    draft_matches = await store.search("服务包", SemanticEmbeddingProvider(), "req-draft", top_k=3, include_draft=True)
    assert draft_matches
    assert draft_matches[0]["id"] == "draft_package"

    rejected_matches = await store.search("保证治好", SemanticEmbeddingProvider(), "req-rejected", top_k=3, include_draft=True)
    assert rejected_matches == []


def test_jiyangjia_query_is_searched_instead_of_rejected_before_retrieval() -> None:
    store = SQLiteKnowledgeStore(":memory:")
    policy = store.classify_query("这个项目能治病吗")
    assert policy == {"action": "search", "scope": "jiyangjia", "category": None, "message": None}


def test_kiosk_short_queries_expand_without_changing_unrelated_questions() -> None:
    store = SQLiteKnowledgeStore(":memory:")

    assert store.classify_query("怎么体验？")["scope"] == "jiyangjia"
    assert store.classify_query("这个游戏怎么体验？")["scope"] == "general"
    assert "智能产品是什么" in _expand_search_query("有什么产品？")
    assert "第一次来怎么办" in _expand_search_query("怎么体验？")
    assert _expand_search_query("积养家是什么？") == "积养家是什么？"


def test_general_query_bypasses_business_knowledge_search() -> None:
    asyncio.run(_assert_general_query_bypasses_business_knowledge_search())


async def _assert_general_query_bypasses_business_knowledge_search() -> None:
    store = SQLiteKnowledgeStore(":memory:")
    await store.index_documents(_documents(), SemanticEmbeddingProvider(), "req-index")

    policy = store.classify_query("今天天气怎么样？")
    matches = await store.search("今天天气怎么样？", SemanticEmbeddingProvider(), "req-weather")

    assert policy["scope"] == "general"
    assert matches == []


def test_unknown_jiyangjia_product_does_not_use_a_weak_semantic_match() -> None:
    asyncio.run(_assert_unknown_jiyangjia_product_does_not_use_a_weak_semantic_match())


async def _assert_unknown_jiyangjia_product_does_not_use_a_weak_semantic_match() -> None:
    store = SQLiteKnowledgeStore(":memory:")
    await store.index_documents(_documents(), SemanticEmbeddingProvider(), "req-index")

    matches = await store.search("积养家有榴莲吗？", SemanticEmbeddingProvider(), "req-unknown")

    assert matches == []


@pytest.mark.parametrize(
    ("query", "expected_id"),
    [
        ("积养家会员余额怎么查询？", "approved_member_boundary"),
        ("帮我查一下积养家会员余额。", "approved_member_boundary"),
        ("积养家今天价格是多少？", "approved_price_boundary"),
    ],
)
def test_brand_prefix_does_not_dilute_approved_boundary_matches(query: str, expected_id: str) -> None:
    asyncio.run(_assert_brand_prefix_does_not_dilute_approved_boundary_matches(query, expected_id))


async def _assert_brand_prefix_does_not_dilute_approved_boundary_matches(query: str, expected_id: str) -> None:
    store = SQLiteKnowledgeStore(":memory:")
    store.upsert_many(
        [
            KnowledgeDocument(
                "approved_member_boundary",
                "会员余额查询",
                "会员余额和账户信息属于个人隐私，我不能查询。请到前台或登录小程序查看。",
                "approved",
            ),
            KnowledgeDocument(
                "approved_price_boundary",
                "问价格转人工",
                "价格是多少，多少钱，怎么收费，请咨询现场工作人员或查看价目表。",
                "approved",
            ),
        ]
    )

    matches = await store.search(query, top_k=3)

    assert matches
    assert matches[0]["id"] == expected_id


def test_parse_selected_json_and_markdown(tmp_path: Path) -> None:
    docs = parse_selected_knowledge_path(Path("knowledge-test/faq_mvp_approved.example.json"))
    assert len(docs) >= 10
    assert {doc.status for doc in docs} >= {"approved", "draft", "rejected"}

    markdown = tmp_path / "selected.md"
    markdown.write_text("# 已选资料\n\n只用于测试解析。", encoding="utf-8")
    parsed = parse_selected_knowledge_path(markdown, default_status="draft")
    assert parsed[0].title == "已选资料"
    assert parsed[0].status == "draft"
    assert parsed[0].source_type == "md"

    text_file = tmp_path / "selected.txt"
    text_file.write_text("TXT 已选资料\n\n只用于测试解析。", encoding="utf-8")
    parsed_text = parse_selected_knowledge_path(text_file, default_status="draft")
    assert parsed_text[0].source_type == "txt"

    from docx import Document

    docx_file = tmp_path / "selected.docx"
    document = Document()
    document.add_paragraph("DOCX 已选资料")
    document.add_paragraph("只用于测试解析。")
    document.save(docx_file)
    parsed_docx = parse_selected_knowledge_path(docx_file, default_status="draft")
    assert parsed_docx[0].source_type == "docx"
    assert "DOCX 已选资料" in parsed_docx[0].text


def _documents() -> list[KnowledgeDocument]:
    return [
        KnowledgeDocument("approved_time", "服务时间", "门店服务时间请以当天现场公告为准。", "approved", "manual://time"),
        KnowledgeDocument("approved_audio", "语音咨询", "点击开始咨询后可以语音提问，并查看字幕。", "approved", "manual://audio"),
        KnowledgeDocument("draft_package", "待确认服务包", "服务包内容仍在待确认，不能正式回答。", "draft", "manual://draft"),
        KnowledgeDocument("rejected_claim", "禁止疗效承诺", "禁止承诺保证治好。", "rejected", "manual://rejected"),
    ]


def _vector_for_text(text: str) -> list[float]:
    if "服务时间" in text or "营业" in text or "现场公告" in text:
        return [1.0, 0.0, 0.0, 0.0]
    if "语音" in text or "字幕" in text:
        return [0.0, 1.0, 0.0, 0.0]
    if "服务包" in text:
        return [0.0, 0.0, 1.0, 0.0]
    return [0.0, 0.0, 0.0, 1.0]
